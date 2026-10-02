"""Format-dispatching document loading for the ingestion pipeline.

``load_document(path)`` returns a normalized :class:`LoadedDocument` no matter
which of the supported formats the file uses. All format-specific parsing
(PDF / DOCX / HTML) lives in this module, so the indexing layer never needs to
know how a document was encoded.

The set of supported extensions is centralised in :data:`SUPPORTED_EXTENSIONS`;
the indexer imports that set instead of maintaining its own copy.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from app.ingestion.frontmatter import (
    default_metadata,
    parse_front_matter,
    remove_front_matter,
)


SUPPORTED_EXTENSIONS: set[str] = {
    ".md",
    ".txt",
    ".pdf",
    ".docx",
    ".html",
    ".htm",
}

FILE_TYPE_MARKDOWN = "markdown"
FILE_TYPE_TEXT = "text"
FILE_TYPE_PDF = "pdf"
FILE_TYPE_DOCX = "docx"
FILE_TYPE_HTML = "html"

_EXTENSION_FILE_TYPES = {
    ".md": FILE_TYPE_MARKDOWN,
    ".txt": FILE_TYPE_TEXT,
    ".pdf": FILE_TYPE_PDF,
    ".docx": FILE_TYPE_DOCX,
    ".html": FILE_TYPE_HTML,
    ".htm": FILE_TYPE_HTML,
}


class DocumentLoadError(ValueError):
    """Raised when a document exists but cannot be read or parsed."""


class UnsupportedFormatError(DocumentLoadError):
    """Raised for file extensions outside :data:`SUPPORTED_EXTENSIONS`."""


class OCRNotSupportedError(DocumentLoadError):
    """Raised when a PDF has no extractable text layer (scan/image only)."""


@dataclass
class LoadedDocument:
    """Format-independent representation of an ingested document.

    ``text`` is the normalized visible text with any Markdown front matter
    already removed. ``heading_aware`` is ``True`` when ``text`` contains
    Markdown-style headings that the heading-aware chunker can split on.
    ``pages`` is populated for PDFs only, keeping one entry per page.
    """

    text: str
    source: str
    document_type: str = "guide"
    metadata: dict = field(default_factory=dict)
    file_type: str = FILE_TYPE_TEXT
    heading_aware: bool = False
    pages: list[str] | None = None
    raw_text: str | None = None


class DocumentLoader:
    """Base class for the per-format loaders."""

    file_type: str = FILE_TYPE_TEXT
    heading_aware: bool = False
    extensions: tuple[str, ...] = ()

    def load(self, path: Path) -> LoadedDocument:
        raise NotImplementedError

    def _read_text(self, path: Path) -> str:
        try:
            return path.read_text(encoding="utf-8")
        except UnicodeDecodeError as error:
            raise DocumentLoadError(
                f"'{path.name}' is not valid UTF-8 text: {error}"
            ) from error

    def _build_metadata(self, file_type: str) -> dict:
        metadata = default_metadata()
        metadata["file_type"] = file_type
        return metadata


class MarkdownLoader(DocumentLoader):
    """Markdown with optional YAML front matter (previous behaviour)."""

    file_type = FILE_TYPE_MARKDOWN
    heading_aware = True
    extensions = (".md",)

    def load(self, path: Path) -> LoadedDocument:

        raw_text = self._read_text(path)

        metadata = parse_front_matter(raw_text)
        metadata["file_type"] = self.file_type

        return LoadedDocument(
            text=remove_front_matter(raw_text),
            source=path.name,
            document_type=metadata["document_type"],
            metadata=metadata,
            file_type=self.file_type,
            heading_aware=True,
            raw_text=raw_text,
        )


class TextLoader(DocumentLoader):
    """Plain UTF-8 text with generic chunking."""

    file_type = FILE_TYPE_TEXT
    heading_aware = False
    extensions = (".txt",)

    def load(self, path: Path) -> LoadedDocument:

        raw_text = self._read_text(path)

        metadata = self._build_metadata(self.file_type)

        return LoadedDocument(
            text=raw_text,
            source=path.name,
            document_type=metadata["document_type"],
            metadata=metadata,
            file_type=self.file_type,
            heading_aware=False,
            raw_text=raw_text,
        )


class PDFLoader(DocumentLoader):
    """PDF text extraction via pypdf (no OCR, no images)."""

    file_type = FILE_TYPE_PDF
    heading_aware = False
    extensions = (".pdf",)

    def load(self, path: Path) -> LoadedDocument:

        try:
            from pypdf import PdfReader
            from pypdf.errors import PdfReadError
        except ImportError as error:  # pragma: no cover - dependency guard
            raise DocumentLoadError(
                "PDF support requires the 'pypdf' package."
            ) from error

        try:
            reader = PdfReader(str(path))
        except PdfReadError as error:
            raise DocumentLoadError(
                f"'{path.name}' is not a valid PDF: {error}"
            ) from error
        except Exception as error:
            raise DocumentLoadError(
                f"could not read PDF '{path.name}': {error}"
            ) from error

        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as error:
                raise DocumentLoadError(
                    f"'{path.name}' is encrypted and cannot be read."
                ) from error

        pages: list[str] = []

        for page_number, page in enumerate(reader.pages, start=1):

            try:
                page_text = page.extract_text() or ""
            except Exception as error:
                raise DocumentLoadError(
                    f"failed to extract text from page "
                    f"{page_number} of '{path.name}': {error}"
                ) from error

            pages.append(page_text)

        if not any(page.strip() for page in pages):
            raise OCRNotSupportedError(
                f"'{path.name}' contains no extractable text. Scanned "
                f"PDFs and OCR are not supported in this phase."
            )

        metadata = self._build_metadata(self.file_type)
        metadata["page_count"] = len(pages)

        text = "\n\n".join(
            page.strip()
            for page in pages
            if page.strip()
        )

        return LoadedDocument(
            text=text,
            source=path.name,
            document_type=metadata["document_type"],
            metadata=metadata,
            file_type=self.file_type,
            heading_aware=False,
            pages=pages,
        )


_HEADING_STYLE_RE = re.compile(r"heading\s+(\d+)")


def _docx_heading_level(style_name: str) -> int:
    """Map a DOCX paragraph style name to a Markdown heading level."""

    name = (style_name or "").strip().lower()

    if name == "title":
        return 1

    match = _HEADING_STYLE_RE.match(name)

    if match:
        return int(match.group(1))

    return 0


class DocxLoader(DocumentLoader):
    """Word documents via python-docx. Paragraph text only."""

    file_type = FILE_TYPE_DOCX
    heading_aware = True
    extensions = (".docx",)

    def load(self, path: Path) -> LoadedDocument:

        try:
            from docx import Document

            document = Document(str(path))
        except ImportError as error:  # pragma: no cover - dependency guard
            raise DocumentLoadError(
                "DOCX support requires the 'python-docx' package."
            ) from error
        except Exception as error:
            raise DocumentLoadError(
                f"'{path.name}' is not a valid DOCX file: {error}"
            ) from error

        lines: list[str] = []
        has_heading = False

        # ``document.paragraphs`` preserves document order. Embedded images
        # and tables are deliberately ignored in this phase.
        for paragraph in document.paragraphs:

            text = (paragraph.text or "").strip()

            if not text:
                continue

            style_name = (
                paragraph.style.name
                if paragraph.style is not None
                else ""
            )

            level = _docx_heading_level(style_name)

            if level:
                has_heading = True
                lines.append(
                    f"{'#' * min(level, 3)} {text}"
                )
            else:
                lines.append(text)

        metadata = self._build_metadata(self.file_type)

        return LoadedDocument(
            text="\n".join(lines),
            source=path.name,
            document_type=metadata["document_type"],
            metadata=metadata,
            file_type=self.file_type,
            heading_aware=has_heading,
        )


_HEADING_TAGS = {
    "h1": 1,
    "h2": 2,
    "h3": 3,
    "h4": 4,
    "h5": 5,
    "h6": 6,
}

_BLOCK_TAGS = (
    "h1",
    "h2",
    "h3",
    "h4",
    "h5",
    "h6",
    "p",
    "li",
    "blockquote",
    "pre",
    "dt",
    "dd",
    "td",
    "th",
)

_STRIPPED_TAGS = ("script", "style", "noscript")


class HTMLLoader(DocumentLoader):
    """HTML visible text via BeautifulSoup. Scripts are never executed."""

    file_type = FILE_TYPE_HTML
    heading_aware = True
    extensions = (".html", ".htm")

    def load(self, path: Path) -> LoadedDocument:

        try:
            from bs4 import BeautifulSoup
        except ImportError as error:  # pragma: no cover - dependency guard
            raise DocumentLoadError(
                "HTML support requires the 'beautifulsoup4' package."
            ) from error

        try:
            raw_text = path.read_text(encoding="utf-8", errors="replace")
        except Exception as error:
            raise DocumentLoadError(
                f"could not read HTML file '{path.name}': {error}"
            ) from error

        if not raw_text.strip():
            raise DocumentLoadError(
                f"'{path.name}' is empty."
            )

        try:
            soup = BeautifulSoup(raw_text, "html.parser")
        except Exception as error:
            raise DocumentLoadError(
                f"'{path.name}' is not valid HTML: {error}"
            ) from error

        # Never let scripts/styles become searchable content.
        for tag in soup(_STRIPPED_TAGS):
            tag.decompose()

        root = soup.body if soup.body is not None else soup

        lines: list[str] = []

        for element in root.find_all(_BLOCK_TAGS):

            # Skip containers so nested blocks are not emitted twice.
            if element.find(_BLOCK_TAGS) is not None:
                continue

            text = element.get_text(" ", strip=True)

            if not text:
                continue

            name = element.name.lower()

            if name in _HEADING_TAGS:
                level = min(_HEADING_TAGS[name], 3)
                lines.append(f"{'#' * level} {text}")
            elif name == "li":
                lines.append(f"- {text}")
            else:
                lines.append(text)

        if not lines:
            # Fragment/plain-text HTML with no block elements.
            fallback = root.get_text("\n", strip=True)

            lines = [
                line
                for line in (chunk.strip() for chunk in fallback.splitlines())
                if line
            ]

        if not lines:
            raise DocumentLoadError(
                f"'{path.name}' contains no visible text."
            )

        metadata = self._build_metadata(self.file_type)

        return LoadedDocument(
            text="\n".join(lines),
            source=path.name,
            document_type=metadata["document_type"],
            metadata=metadata,
            file_type=self.file_type,
            heading_aware=any(
                line.startswith("#")
                for line in lines
            ),
        )


_LOADERS: dict[str, DocumentLoader] = {}


def _register(loader: DocumentLoader) -> None:
    for extension in loader.extensions:
        _LOADERS[extension] = loader


_register(MarkdownLoader())
_register(TextLoader())
_register(PDFLoader())
_register(DocxLoader())
_register(HTMLLoader())


def is_supported(path: str | Path) -> bool:
    """Return ``True`` when the path's extension is supported."""

    return Path(path).suffix.lower() in SUPPORTED_EXTENSIONS


def get_loader(path: str | Path) -> DocumentLoader:
    """Return the loader responsible for a path's extension."""

    suffix = Path(path).suffix.lower()

    loader = _LOADERS.get(suffix)

    if loader is None:
        raise UnsupportedFormatError(
            f"Unsupported document type: {suffix or '<none>'}"
        )

    return loader


def load_document(path: str) -> LoadedDocument:
    """Load any supported document into a normalized :class:`LoadedDocument`."""

    file_path = Path(path)

    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    return get_loader(file_path).load(file_path)


def load_text_file(path: str) -> str:
    """Read a plain-text or Markdown file as raw UTF-8 text.

    Kept for the existing local experimentation scripts.
    """

    file_path = Path(path)

    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    if file_path.suffix.lower() not in {".txt", ".md"}:
        raise UnsupportedFormatError(
            f"Unsupported text file type: {file_path.suffix}"
        )

    return file_path.read_text(encoding="utf-8")
