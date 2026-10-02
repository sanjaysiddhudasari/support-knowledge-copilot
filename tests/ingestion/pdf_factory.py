"""Helpers that build small real PDF/DOCX files for loader tests.

The PDF writer emits a minimal, uncompressed, valid PDF so tests do not need
a heavyweight generation dependency (no reportlab/fpdf). Only standard fonts
and simple ``Tj`` text operators are used, which pypdf extracts reliably.
"""

from pathlib import Path


def build_pdf_bytes(pages: list[str], include_text: bool = True) -> bytes:
    """Return a minimal valid PDF with one page per entry in ``pages``.

    When ``include_text`` is ``False`` the content streams are empty, which
    emulates a scanned/image-only PDF with no extractable text layer.
    """

    # Object numbering: 1 catalog, 2 pages tree, 3 font, then two objects per
    # page (page object, content stream object).
    page_ids = [4 + 2 * index for index in range(len(pages))]
    content_ids = [5 + 2 * index for index in range(len(pages))]

    objects: dict[int, str] = {}

    objects[1] = "<< /Type /Catalog /Pages 2 0 R >>"

    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)

    objects[2] = (
        f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>"
    )

    objects[3] = (
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    )

    for page_id, content_id, text in zip(page_ids, content_ids, pages):

        objects[page_id] = (
            "<< /Type /Page /Parent 2 0 R "
            "/MediaBox [0 0 612 792] "
            "/Resources << /Font << /F1 3 0 R >> >> "
            f"/Contents {content_id} 0 R >>"
        )

        if include_text:
            escaped = (
                text.replace("\\", "\\\\")
                .replace("(", "\\(")
                .replace(")", "\\)")
            )
            stream = f"BT /F1 14 Tf 72 720 Td ({escaped}) Tj ET"
        else:
            stream = ""

        objects[content_id] = (
            f"<< /Length {len(stream)} >>\n"
            f"stream\n{stream}\nendstream"
        )

    out = bytearray(b"%PDF-1.4\n")

    offsets: dict[int, int] = {}

    for number in sorted(objects):
        offsets[number] = len(out)
        out += f"{number} 0 obj\n{objects[number]}\nendobj\n".encode(
            "latin-1"
        )

    xref_offset = len(out)
    size = max(objects) + 1

    out += f"xref\n0 {size}\n".encode("latin-1")
    out += b"0000000000 65535 f \n"

    for number in range(1, size):
        out += f"{offsets[number]:010d} 00000 n \n".encode("latin-1")

    out += (
        f"trailer\n<< /Size {size} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n"
    ).encode("latin-1")

    return bytes(out)


def write_pdf(path: Path, pages: list[str], include_text: bool = True) -> Path:
    path.write_bytes(build_pdf_bytes(pages, include_text=include_text))
    return path


def write_docx(path: Path, paragraphs: list[tuple[str, str]]) -> Path:
    """Write a DOCX from ``(kind, text)`` tuples.

    ``kind`` is one of ``"heading1"``, ``"heading2"`` or ``"paragraph"``.
    """

    from docx import Document

    document = Document()

    for kind, text in paragraphs:

        if kind == "heading1":
            document.add_heading(text, level=1)
        elif kind == "heading2":
            document.add_heading(text, level=2)
        else:
            document.add_paragraph(text)

    document.save(str(path))

    return path
